import pandas as pd

            #一、从零开始，创建一个全新的 Excel 文件
# 1. 准备你要写入的数据（用字典格式，键是列名，值是列表）
data = {
    "姓名": ["王五", "赵六"],
    "年龄": [22, 28],
    "城市": ["北京", "深圳"],
    "销售额": [500.0, 800.0],
    "订单状态": ["已完成", "处理中"]
}

# 2. 将数据转换成 Pandas 的 DataFrame 结构
df_new = pd.DataFrame(data)

# 3. 写入到一个全新的 Excel 文件中
# index=False 表示不需要把左侧的 0, 1, 2 这样的行索引写进 Excel
file_path = r"E:\study\python\data\全新的表格.xlsx"
df_new.to_excel(file_path, index=False)

print("全新的 Excel 文件创建并写入成功！")

            #在已有的 Excel 文件中“插入/追加”新信息
#1.在表格末尾直接“追加”新行（最常用）
from openpyxl import load_workbook

# 1. 打开已经存在的 Excel 文件
file_path = r"E:\study\python\data\学生表.xlsx"
wb = load_workbook(file_path)

# 2. 激活当前正在使用的这页工作表（Sheet）
ws = wb.active

# 3. 准备要插入的新行数据（用列表表示）
new_row = ["孙悟空", 500, "花果山", 9999.0, "已完成"]

# 4. 使用 append() 方法，它会自动找到最后一行的下一行并写进去
ws.append(new_row)

# 5. 一定要保存文件，修改才会生效！
wb.save(file_path)

print("新数据已成功追加到已有表格的末尾！")

#2.精准定位，在指定单元格“修改/插入”数据
from openpyxl import load_workbook

file_path = r"E:\study\python\data\学生表.xlsx"
wb = load_workbook(file_path)
ws = wb.active

# 写法 1：通过 Excel 的坐标定位（比如 B2 单元格）
ws["B2"] = 100

# 写法 2：通过行号和列号定位（比如第 5 行，第 4 列）
ws.cell(row=5, column=4, value=2000.0)

# 保存
wb.save(file_path)
print("指定单元格数据修改成功！")